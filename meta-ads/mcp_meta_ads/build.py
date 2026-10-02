"""Spreadsheet-driven Meta campaign builder.

The operator hands over a folder containing one workbook (`campaign.xlsx`,
three sheets: Campaign / AdSets / Ads) plus the image and video files the
Ads sheet references. This module turns that into a PAUSED campaign on the
ad account, in two phases:

  plan(...)    — parse + validate the workbook, resolve every city / zip /
                 interest / audience / page against Meta, and return the
                 exact payloads that WOULD be sent. Creates nothing. Uploads
                 nothing. This is the review gate.
  execute(...) — upload assets, create campaign → ad sets → creatives → ads,
                 all PAUSED. Resumable: every created ID is written to a state
                 file next to the workbook, so a failure halfway through
                 doesn't duplicate anything on rerun.

Nothing here flips a status to ACTIVE. The operator does that in Ads Manager
after reviewing the build.

Deliberately free of MCP imports so it can be unit-tested without the server.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from . import graph

# ---------------------------------------------------------------------------
# Sheet schema
# ---------------------------------------------------------------------------

CAMPAIGN_KEYS: dict[str, str] = {
    # key: help text (also written into the template)
    "client": "Brand name (used in file names only).",
    "ad_account_id": "act_XXXXXXXXX. Required.",
    "page_id": "Facebook Page ID the ads run as. Required. Use page_list to find it.",
    "instagram_user_id": "Instagram account ID for IG placements (the IG *user* id, e.g. 1784140...). Optional; omit to use the Page-backed IG identity. Legacy header instagram_actor_id is accepted.",
    "campaign_name": "Required. Follow your naming convention.",
    "objective": "OUTCOME_LEADS | OUTCOME_SALES | OUTCOME_TRAFFIC | OUTCOME_AWARENESS | OUTCOME_ENGAGEMENT",
    "buying_type": "AUCTION (default) | RESERVED",
    "special_ad_categories": "Blank for none, or HOUSING | EMPLOYMENT | CREDIT | ISSUES_ELECTIONS_POLITICS (semicolon-separate multiple).",
    "campaign_daily_budget": "Dollars (e.g. 150.00). Fill this OR ad-set budgets, not both. Filling it = CBO (Advantage+ campaign budget).",
    "campaign_lifetime_budget": "Dollars. Requires end_time on every ad set. Mutually exclusive with campaign_daily_budget.",
    "bid_strategy": "LOWEST_COST_WITHOUT_CAP (default) | LOWEST_COST_WITH_BID_CAP | COST_CAP | LOWEST_COST_WITH_MIN_ROAS",
    "pixel_id": "Dataset / Pixel ID. Required when any ad set optimizes for OFFSITE_CONVERSIONS.",
    "conversion_event": "LEAD | PURCHASE | CONTACT | SCHEDULE | COMPLETE_REGISTRATION | SUBMIT_APPLICATION | ... (default LEAD)",
    "default_country": "ISO code used to resolve cities and zips (default US).",
    "url_tags": "Appended to every destination URL. e.g. utm_source=facebook&utm_medium=ppc&utm_campaign={{campaign.name}}&utm_content={{ad.name}}",
}

ADSET_COLUMNS: dict[str, str] = {
    "adset_name": "Required, unique within the campaign.",
    "daily_budget": "Dollars. Leave blank when the campaign uses CBO.",
    "lifetime_budget": "Dollars. Requires end_time. Blank when CBO.",
    "optimization_goal": "OFFSITE_CONVERSIONS | LANDING_PAGE_VIEWS | LINK_CLICKS | REACH | IMPRESSIONS | THRUPLAY | LEAD_GENERATION",
    "billing_event": "IMPRESSIONS (default) | LINK_CLICKS | THRUPLAY",
    "bid_amount": "Dollars. Only for bid cap / cost cap strategies.",
    "geo_type": "countries | radius | cities | zips | regions",
    "geo_value": (
        "countries: US;CA   radius: 123 Main St, Nashville TN 37203 | 25mi ; ...   "
        "cities: Nashville, TN | 15mi ; Franklin, TN   zips: 37203; 37204   regions: Tennessee; Kentucky"
    ),
    "excluded_geo_type": "Optional. Same options as geo_type.",
    "excluded_geo_value": "Optional. Same formats as geo_value.",
    "age_min": "18-65 (default 18).",
    "age_max": "18-65 (default 65).",
    "genders": "all (default) | male | female",
    "interests": "Semicolon-separated interest names. Resolved via Meta search; the plan shows what matched.",
    "custom_audiences": "Semicolon-separated audience names or IDs to INCLUDE.",
    "excluded_audiences": "Semicolon-separated audience names or IDs to EXCLUDE.",
    "advantage_audience": "Y (default when no interests/audiences) | N. Advantage+ audience lets Meta expand beyond your inputs.",
    "placements": "auto (default) or e.g. facebook:feed,story,reels; instagram:stream,story,reels",
    "start_time": "Optional ISO 8601 with offset, e.g. 2026-10-01T06:00:00-0500",
    "end_time": "Optional ISO 8601. Required with lifetime budgets.",
}

AD_COLUMNS: dict[str, str] = {
    "ad_name": "Required, unique within the campaign.",
    "adset_name": "Must match an adset_name on the AdSets sheet.",
    "format": "image | video | carousel | dco",
    "asset_files": "File name(s) in this folder. image/video: one file. carousel: 2-10 files, semicolon-separated, in card order. dco: up to 10 images or videos.",
    "primary_text": "Body copy. For dco: up to 5 variants separated by ||",
    "headline": "For carousel: one per card separated by || (or one to reuse). For dco: up to 5 variants separated by ||",
    "description": "Optional link description. Carousel: per card with ||",
    "destination_url": "Landing page. Carousel: one per card with || (or one to reuse).",
    "display_link": "Optional caption domain shown under the headline (e.g. example.com).",
    "cta": "LEARN_MORE | GET_QUOTE | SIGN_UP | SHOP_NOW | CONTACT_US | GET_OFFER | BOOK_NOW | DOWNLOAD | APPLY_NOW | SUBSCRIBE | WATCH_MORE. dco: multiple with ||",
    "url_tags": "Optional override of the campaign-level url_tags for this ad.",
}

_CTA_ALLOWED = {
    "LEARN_MORE", "GET_QUOTE", "SIGN_UP", "SHOP_NOW", "CONTACT_US", "GET_OFFER", "BOOK_NOW",
    "DOWNLOAD", "APPLY_NOW", "SUBSCRIBE", "WATCH_MORE", "BUY_NOW", "ORDER_NOW", "REQUEST_TIME",
    "SEE_MORE", "SEND_MESSAGE", "WHATSAPP_MESSAGE", "CALL_NOW", "GET_DIRECTIONS", "NO_BUTTON",
}
_OBJECTIVES = {
    "OUTCOME_LEADS", "OUTCOME_SALES", "OUTCOME_TRAFFIC", "OUTCOME_AWARENESS",
    "OUTCOME_ENGAGEMENT", "OUTCOME_APP_PROMOTION",
}
_OPT_GOALS = {
    "OFFSITE_CONVERSIONS", "LANDING_PAGE_VIEWS", "LINK_CLICKS", "REACH", "IMPRESSIONS",
    "THRUPLAY", "LEAD_GENERATION", "POST_ENGAGEMENT", "VALUE", "CONVERSATIONS",
}
# Which performance goals Meta accepts under each ODAX objective (error 2490408
# otherwise). Kept deliberately conservative — the common combinations only.
_GOALS_BY_OBJECTIVE: dict[str, set[str]] = {
    "OUTCOME_SALES": {"OFFSITE_CONVERSIONS", "VALUE", "CONVERSATIONS", "LINK_CLICKS", "LANDING_PAGE_VIEWS", "IMPRESSIONS", "REACH"},
    "OUTCOME_LEADS": {"OFFSITE_CONVERSIONS", "LEAD_GENERATION", "QUALITY_LEAD", "CONVERSATIONS", "LINK_CLICKS", "LANDING_PAGE_VIEWS", "IMPRESSIONS", "REACH"},
    "OUTCOME_TRAFFIC": {"LINK_CLICKS", "LANDING_PAGE_VIEWS", "IMPRESSIONS", "REACH", "CONVERSATIONS"},
    "OUTCOME_ENGAGEMENT": {"POST_ENGAGEMENT", "THRUPLAY", "CONVERSATIONS", "PAGE_LIKES", "EVENT_RESPONSES", "OFFSITE_CONVERSIONS", "LINK_CLICKS", "LANDING_PAGE_VIEWS", "IMPRESSIONS", "REACH"},
    "OUTCOME_AWARENESS": {"REACH", "IMPRESSIONS", "AD_RECALL_LIFT", "THRUPLAY"},
    "OUTCOME_APP_PROMOTION": {"APP_INSTALLS", "OFFSITE_CONVERSIONS", "VALUE", "LINK_CLICKS"},
}
# Observed live on act_1484941429463480: OUTCOME_SALES rejected LANDING_PAGE_VIEWS
# even though Meta's docs list it. Treat the traffic-style goals as sales/leads
# exceptions until proven otherwise.
_GOALS_BY_OBJECTIVE["OUTCOME_SALES"] -= {"LANDING_PAGE_VIEWS", "LINK_CLICKS", "IMPRESSIONS", "REACH"}
_GOALS_BY_OBJECTIVE["OUTCOME_LEADS"] -= {"LANDING_PAGE_VIEWS", "LINK_CLICKS", "IMPRESSIONS", "REACH"}

_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
_VIDEO_EXT = {".mp4", ".mov", ".m4v", ".avi", ".webm"}

_US_STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware", "FL": "Florida", "GA": "Georgia",
    "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts",
    "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana",
    "NE": "Nebraska", "NV": "Nevada", "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico",
    "NY": "New York", "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
    "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina",
    "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas", "UT": "Utah", "VT": "Vermont",
    "VA": "Virginia", "WA": "Washington", "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming",
    "DC": "District of Columbia",
}
_CA_PROVINCES = {
    "AB": "Alberta", "BC": "British Columbia", "MB": "Manitoba", "NB": "New Brunswick",
    "NL": "Newfoundland and Labrador", "NS": "Nova Scotia", "ON": "Ontario",
    "PE": "Prince Edward Island", "QC": "Quebec", "SK": "Saskatchewan",
}


class BuildError(ValueError):
    """A validation problem the operator has to fix in the workbook."""


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------


def _s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _split(v: Any, sep: str = ";") -> list[str]:
    return [p.strip() for p in _s(v).split(sep) if p.strip()]


def _variants(v: Any) -> list[str]:
    """`||`-separated copy variants."""
    return [p.strip() for p in _s(v).split("||") if p.strip()]


def _dollars_to_minor(v: Any, what: str) -> int | None:
    s = _s(v).replace("$", "").replace(",", "")
    if not s:
        return None
    try:
        return int(round(float(s) * 100))
    except ValueError as e:
        raise BuildError(f"{what}: '{v}' is not a dollar amount.") from e


def _yes(v: Any, default: bool) -> bool:
    s = _s(v).lower()
    if not s:
        return default
    return s in {"y", "yes", "true", "1"}


def _int(v: Any, default: int, what: str) -> int:
    s = _s(v)
    if not s:
        return default
    try:
        return int(float(s))
    except ValueError as e:
        raise BuildError(f"{what}: '{v}' is not a whole number.") from e


def _find_csv_trio(folder: Path) -> dict[str, Path]:
    """Google Sheets exports one CSV per tab, named '<Sheet name> - <Tab>.csv'.
    Match each tab by the word at the end of the filename. 'adsets' is checked
    before 'ads' so 'Acme - AdSets.csv' doesn't get claimed as the Ads tab."""
    found: dict[str, Path] = {}
    csvs = sorted(folder.glob("*.csv"))
    for tab in ("campaign", "adsets", "ads"):
        for f in csvs:
            stem = re.sub(r"[\s_\-]+", " ", f.stem.lower()).strip()
            last = stem.split(" ")[-1]
            if last == tab and f not in found.values():
                found[tab] = f
                break
    return found


def _read_csv_rows(p: Path) -> list[list[str]]:
    with p.open(newline="", encoding="utf-8-sig") as fh:
        return [list(r) for r in csv.reader(fh)]


def read_workbook(path: str | Path) -> dict[str, Any]:
    """Return {"campaign": {key: value}, "adsets": [row...], "ads": [row...]}.

    Accepts:
      * an .xlsx with sheets Campaign / AdSets / Ads (Google Sheets → File →
        Download → Microsoft Excel keeps all tabs), or
      * a folder (or any .csv inside it) holding one CSV per tab, as Google
        Sheets produces when each tab is downloaded as CSV.
    """
    p = Path(path)
    if p.is_dir() or p.suffix.lower() == ".csv":
        folder = p if p.is_dir() else p.parent
        trio = _find_csv_trio(folder)
        missing = [t for t in ("campaign", "adsets", "ads") if t not in trio]
        if missing:
            xlsx = sorted(folder.glob("*.xlsx"))
            if xlsx and not p.suffix.lower() == ".csv":
                return read_workbook(xlsx[0])
            raise BuildError(
                f"Could not find CSVs for {missing} in {folder}. Expected filenames ending in "
                f"'Campaign.csv', 'AdSets.csv', 'Ads.csv' (Google Sheets tab export), or a campaign.xlsx."
            )
        campaign: dict[str, str] = {}
        for row in _read_csv_rows(trio["campaign"]):
            if not row or _s(row[0]) == "" or _s(row[0]).lower() == "key":
                continue
            campaign[_s(row[0]).lower()] = _s(row[1]) if len(row) > 1 else ""

        def csv_table(f: Path) -> list[dict[str, str]]:
            rows = _read_csv_rows(f)
            if not rows:
                return []
            header = [_s(h).lower() for h in rows[0]]
            out = []
            for r in rows[1:]:
                if not r or all(_s(c) == "" for c in r):
                    continue
                rec = {header[i]: _s(r[i]) for i in range(min(len(header), len(r))) if header[i]}
                if rec.get(next(iter(rec), ""), "").startswith("#"):
                    continue
                out.append(rec)
            return out

        return {"campaign": campaign, "adsets": csv_table(trio["adsets"]), "ads": csv_table(trio["ads"])}

    from openpyxl import load_workbook

    if not p.is_file():
        raise BuildError(f"Workbook not found: {p}")
    wb = load_workbook(p, data_only=True, read_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    for need in ("campaign", "adsets", "ads"):
        if need not in names:
            raise BuildError(f"Workbook is missing a sheet named '{need}'. Found: {wb.sheetnames}")

    campaign: dict[str, str] = {}
    for row in wb[names["campaign"]].iter_rows(values_only=True):
        if not row or _s(row[0]) == "" or _s(row[0]).lower() == "key":
            continue
        campaign[_s(row[0]).lower()] = _s(row[1]) if len(row) > 1 else ""

    def table(sheet_name: str) -> list[dict[str, str]]:
        rows = list(wb[names[sheet_name]].iter_rows(values_only=True))
        if not rows:
            return []
        header = [_s(h).lower() for h in rows[0]]
        out = []
        for r in rows[1:]:
            if not r or all(_s(c) == "" for c in r):
                continue
            rec = {header[i]: _s(r[i]) for i in range(min(len(header), len(r))) if header[i]}
            # Skip help/example rows the template ships with.
            if rec.get(next(iter(rec), ""), "").startswith("#"):
                continue
            out.append(rec)
        return out

    return {"campaign": campaign, "adsets": table("adsets"), "ads": table("ads")}


# ---------------------------------------------------------------------------
# Resolver — the only piece that talks to Meta during planning
# ---------------------------------------------------------------------------


@dataclass
class Resolver:
    """Resolves human inputs (city names, interest names, audience names) to
    Meta IDs/keys, with per-run caching. Swappable in tests."""

    ad_account_id: str
    default_country: str = "US"
    notes: list[str] = field(default_factory=list)
    _audiences: list[dict] | None = None
    _cache: dict[str, Any] = field(default_factory=dict)

    # -- geo -------------------------------------------------------------

    def _region_name(self, abbr_or_name: str, country: str) -> str:
        table = _US_STATES if country == "US" else _CA_PROVINCES if country == "CA" else {}
        return table.get(abbr_or_name.upper(), abbr_or_name)

    def city(self, spec: str) -> dict[str, Any]:
        """'Nashville, TN | 15mi' → {key, radius, distance_unit}."""
        key = f"city:{spec}"
        if key in self._cache:
            return self._cache[key]
        name_part, _, radius_part = spec.partition("|")
        city, _, region = [x.strip() for x in name_part.partition(",")]
        country = self.default_country
        if region and len(region) == 2 and region.upper() not in _US_STATES and region.upper() in _CA_PROVINCES:
            country = "CA"
        rows = graph.search_geo(city, ["city"], country_code=country)
        want_region = self._region_name(region, country).lower() if region else None
        match = None
        for r in rows:
            if r.get("name", "").lower() != city.lower():
                continue
            if want_region and (r.get("region") or "").lower() != want_region:
                continue
            match = r
            break
        if match is None and rows and not want_region:
            match = rows[0]
        if match is None:
            options = "; ".join(f"{r.get('name')}, {r.get('region')}" for r in rows[:6]) or "none"
            raise BuildError(f"City '{spec}' did not resolve. Meta suggestions: {options}")
        out: dict[str, Any] = {"key": match["key"], "name": f"{match.get('name')}, {match.get('region')}"}
        if radius_part.strip():
            miles = _parse_miles(radius_part)
            if not 10 <= miles <= 50:
                raise BuildError(f"City radius for '{spec}' must be 10-50 miles (Meta limit).")
            out["radius"] = miles
            out["distance_unit"] = "mile"
        self.notes.append(f"city '{spec}' → {out['name']} ({out['key']})")
        self._cache[key] = out
        return out

    def zip(self, code: str) -> dict[str, Any]:
        key = f"zip:{code}"
        if key in self._cache:
            return self._cache[key]
        rows = graph.search_geo(code, ["zip"], country_code=self.default_country)
        match = next((r for r in rows if r.get("name") == code or r.get("key", "").endswith(f":{code}")), None)
        if match is None:
            raise BuildError(f"Zip '{code}' did not resolve in {self.default_country}.")
        out = {"key": match["key"], "name": f"{code} ({match.get('primary_city') or match.get('region')})"}
        self._cache[key] = out
        return out

    def region(self, name: str) -> dict[str, Any]:
        key = f"region:{name}"
        if key in self._cache:
            return self._cache[key]
        full = self._region_name(name, self.default_country)
        rows = graph.search_geo(full, ["region"], country_code=self.default_country)
        match = next((r for r in rows if r.get("name", "").lower() == full.lower()), rows[0] if rows else None)
        if match is None:
            raise BuildError(f"Region '{name}' did not resolve.")
        out = {"key": match["key"], "name": match.get("name")}
        self._cache[key] = out
        return out

    # -- interests / audiences --------------------------------------------

    def interest(self, name: str) -> dict[str, Any]:
        key = f"interest:{name.lower()}"
        if key in self._cache:
            return self._cache[key]
        rows = graph.search_interests(name)
        exact = next((r for r in rows if r.get("name", "").lower() == name.lower()), None)
        match = exact or (rows[0] if rows else None)
        if match is None:
            raise BuildError(f"Interest '{name}' returned no results from Meta.")
        if not exact:
            self.notes.append(f"interest '{name}' → nearest match '{match.get('name')}' (verify)")
        else:
            size = match.get("audience_size_lower_bound")
            self.notes.append(f"interest '{name}' → {match['id']}" + (f" (~{size:,}+ people)" if size else ""))
        out = {"id": match["id"], "name": match.get("name")}
        self._cache[key] = out
        return out

    def audience(self, name_or_id: str) -> dict[str, Any]:
        if self._audiences is None:
            self._audiences = graph.list_custom_audiences(self.ad_account_id)
        if re.fullmatch(r"\d{6,}", name_or_id):
            match = next((a for a in self._audiences if a.get("id") == name_or_id), {"id": name_or_id, "name": name_or_id})
        else:
            match = next((a for a in self._audiences if a.get("name", "").lower() == name_or_id.lower()), None)
        if match is None:
            close = [a["name"] for a in self._audiences if name_or_id.lower() in a.get("name", "").lower()][:6]
            raise BuildError(
                f"Audience '{name_or_id}' not found on {self.ad_account_id}."
                + (f" Similar: {'; '.join(close)}" if close else "")
            )
        self.notes.append(f"audience '{name_or_id}' → {match.get('name')} ({match['id']})")
        return {"id": match["id"], "name": match.get("name")}


def _parse_miles(text: str) -> float:
    m = re.search(r"([\d.]+)\s*(mi|mile|miles|km)?", text.strip().lower())
    if not m:
        raise BuildError(f"Could not read a radius from '{text}' (use e.g. '25mi').")
    val = float(m.group(1))
    if m.group(2) == "km":
        val = val / 1.609
    return round(val, 1)


# ---------------------------------------------------------------------------
# Targeting builder
# ---------------------------------------------------------------------------


def _geo_block(geo_type: str, geo_value: str, resolver: Resolver) -> dict[str, Any]:
    t = geo_type.strip().lower()
    items = _split(geo_value)
    if not t or not items:
        raise BuildError("geo_type and geo_value are both required on every ad set.")
    block: dict[str, Any] = {}
    if t == "countries":
        block["countries"] = [c.upper() for c in items]
    elif t == "radius":
        locs = []
        for item in items:
            addr, _, rad = item.partition("|")
            if not rad.strip():
                raise BuildError(f"radius entry '{item}' needs '| <miles>mi'.")
            miles = _parse_miles(rad)
            if not 1 <= miles <= 50:
                raise BuildError(f"Radius for '{addr.strip()}' must be 1-50 miles.")
            loc: dict[str, Any] = {"radius": miles, "distance_unit": "mile"}
            latlng = re.fullmatch(r"\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*", addr)
            if latlng:
                loc["latitude"], loc["longitude"] = float(latlng.group(1)), float(latlng.group(2))
            else:
                # Meta's geocoder (error 1487855) rejects bare "City, ST" strings;
                # it wants a street address. Catch that before the build fails.
                if not re.search(r"\d", addr):
                    raise BuildError(
                        f"radius entry '{addr.strip()}' has no street number — Meta cannot geocode a "
                        f"bare city name. Use a full street address, 'lat,lng | Nmi', or geo_type=cities."
                    )
                loc["address_string"] = addr.strip()
            locs.append(loc)
        block["custom_locations"] = locs
    elif t == "cities":
        block["cities"] = [resolver.city(i) for i in items]
    elif t == "zips":
        block["zips"] = [{"key": resolver.zip(i)["key"]} for i in items]
    elif t == "regions":
        block["regions"] = [{"key": resolver.region(i)["key"]} for i in items]
    else:
        raise BuildError(f"geo_type '{geo_type}' is not one of countries | radius | cities | zips | regions.")
    # Drop the display-only "name" we tucked into city entries.
    for c in block.get("cities", []):
        c.pop("name", None)
    block["location_types"] = ["home", "recent"]
    return block


def _placements(spec: str) -> dict[str, Any]:
    """'facebook:feed,story,reels; instagram:stream,reels' → publisher_platforms + positions."""
    s = spec.strip().lower()
    if not s or s == "auto":
        return {}
    out: dict[str, Any] = {"publisher_platforms": []}
    for chunk in _split(s):
        platform, _, positions = chunk.partition(":")
        platform = platform.strip()
        if platform not in {"facebook", "instagram", "audience_network", "messenger", "threads"}:
            raise BuildError(f"placements: unknown platform '{platform}'.")
        out["publisher_platforms"].append(platform)
        pos = [p.strip() for p in positions.split(",") if p.strip()]
        if pos and platform in {"facebook", "instagram", "messenger", "audience_network"}:
            out[f"{platform}_positions"] = pos
    return out


def build_targeting(row: dict[str, str], resolver: Resolver) -> dict[str, Any]:
    t: dict[str, Any] = {}
    t["geo_locations"] = _geo_block(row.get("geo_type", ""), row.get("geo_value", ""), resolver)
    if _s(row.get("excluded_geo_type")) and _s(row.get("excluded_geo_value")):
        ex = _geo_block(row["excluded_geo_type"], row["excluded_geo_value"], resolver)
        ex.pop("location_types", None)
        t["excluded_geo_locations"] = ex

    age_min = _int(row.get("age_min"), 18, "age_min")
    age_max = _int(row.get("age_max"), 65, "age_max")
    if not 18 <= age_min <= age_max <= 65:
        raise BuildError(f"age range {age_min}-{age_max} is invalid (18-65, min ≤ max).")
    t["age_min"], t["age_max"] = age_min, age_max

    g = _s(row.get("genders")).lower()
    if g in {"male", "men", "m"}:
        t["genders"] = [1]
    elif g in {"female", "women", "f"}:
        t["genders"] = [2]
    elif g and g != "all":
        raise BuildError(f"genders '{row.get('genders')}' must be all | male | female.")

    interests = [resolver.interest(i) for i in _split(row.get("interests"))]
    if interests:
        t["flexible_spec"] = [{"interests": interests}]

    inc = [resolver.audience(a) for a in _split(row.get("custom_audiences"))]
    exc = [resolver.audience(a) for a in _split(row.get("excluded_audiences"))]
    if inc:
        t["custom_audiences"] = [{"id": a["id"]} for a in inc]
    if exc:
        t["excluded_custom_audiences"] = [{"id": a["id"]} for a in exc]

    has_narrowing = bool(interests or inc)
    adv = _yes(row.get("advantage_audience"), default=not has_narrowing)
    t["targeting_automation"] = {"advantage_audience": 1 if adv else 0}
    if adv and age_max < 65:
        # Meta rule (error 1870189): with Advantage+ audience on, age_max is an
        # audience *control* and must be 65. Lower ceilings are only accepted
        # as suggestions, which the sheet doesn't model. Lift it and say so.
        t["age_max"] = 65
        resolver.notes.append(
            f"age_max {age_max} → 65: Advantage+ audience requires max age 65 "
            f"(set advantage_audience=N to keep a hard {age_max} ceiling)"
        )

    t.update(_placements(_s(row.get("placements"))))
    return t


# ---------------------------------------------------------------------------
# Plan
# ---------------------------------------------------------------------------


@dataclass
class Plan:
    workbook: Path
    folder: Path
    ad_account_id: str
    campaign: dict[str, Any]
    adsets: list[dict[str, Any]]           # {name, payload, row}
    ads: list[dict[str, Any]]              # {name, adset_name, format, assets, creative_payload_fn, ...}
    assets: dict[str, Path]                # relative name → absolute path
    notes: list[str]
    warnings: list[str]

    def summary(self) -> str:
        cbo = "CBO" if self.campaign.get("daily_budget") or self.campaign.get("lifetime_budget") else "ad-set budgets"
        lines = [
            f"# Build plan — {self.campaign['name']}",
            "",
            f"- Account: {self.ad_account_id}   Page: {self.campaign['_page_id']}",
            f"- Objective: {self.campaign['objective']}   Budget mode: {cbo}",
            f"- Ad sets: {len(self.adsets)}   Ads: {len(self.ads)}   Assets: {len(self.assets)}",
            f"- Everything will be created PAUSED.",
            "",
            "## Ad sets",
        ]
        for a in self.adsets:
            p = a["payload"]
            tg = p["targeting"]
            geo = tg["geo_locations"]
            geo_desc = next((f"{k}: {len(v)}" for k, v in geo.items() if k != "location_types"), "?")
            budget = (
                f"${p['daily_budget']/100:,.2f}/day" if p.get("daily_budget")
                else f"${p['lifetime_budget']/100:,.2f} lifetime" if p.get("lifetime_budget")
                else "campaign budget"
            )
            extras = []
            if tg.get("flexible_spec"):
                extras.append(f"{len(tg['flexible_spec'][0]['interests'])} interests")
            if tg.get("custom_audiences"):
                extras.append(f"{len(tg['custom_audiences'])} audiences")
            if tg.get("excluded_custom_audiences"):
                extras.append(f"{len(tg['excluded_custom_audiences'])} excluded")
            extras.append("Advantage+ audience " + ("on" if tg["targeting_automation"]["advantage_audience"] else "off"))
            lines.append(
                f"- **{a['name']}** — {p['optimization_goal']}, {budget}, geo {geo_desc}, "
                f"age {tg['age_min']}-{tg['age_max']}, {', '.join(extras)}"
            )
        lines += ["", "## Ads"]
        for ad in self.ads:
            lines.append(f"- **{ad['name']}** ({ad['format']}) in {ad['adset_name']} — {', '.join(ad['assets'])} → {ad['row'].get('destination_url','')}")
        if self.notes:
            lines += ["", "## Resolutions (verify these)"] + [f"- {n}" for n in self.notes]
        if self.warnings:
            lines += ["", "## Warnings"] + [f"- {w}" for w in self.warnings]
        return "\n".join(lines)


def _find_assets(folder: Path) -> dict[str, Path]:
    out: dict[str, Path] = {}
    for p in folder.rglob("*"):
        if p.is_file() and p.suffix.lower() in _IMAGE_EXT | _VIDEO_EXT and not p.name.startswith("."):
            out.setdefault(p.name, p)
    return out


def plan(workbook_path: str | Path, resolver_factory: Callable[[str, str], Resolver] | None = None) -> Plan:
    wb_path = Path(workbook_path).expanduser().resolve()
    data = read_workbook(wb_path)
    c = data["campaign"]
    warnings: list[str] = []

    # ---- campaign --------------------------------------------------------
    for req in ("ad_account_id", "page_id", "campaign_name", "objective"):
        if not c.get(req):
            raise BuildError(f"Campaign sheet: '{req}' is required.")
    acct = graph.normalize_account_id(c["ad_account_id"])
    objective = c["objective"].upper()
    if objective not in _OBJECTIVES:
        raise BuildError(f"objective '{c['objective']}' must be one of {sorted(_OBJECTIVES)}.")
    daily = _dollars_to_minor(c.get("campaign_daily_budget"), "campaign_daily_budget")
    lifetime = _dollars_to_minor(c.get("campaign_lifetime_budget"), "campaign_lifetime_budget")
    if daily and lifetime:
        raise BuildError("Set campaign_daily_budget OR campaign_lifetime_budget, not both.")
    cbo = bool(daily or lifetime)
    special = [x.upper() for x in _split(c.get("special_ad_categories"))]
    campaign_payload: dict[str, Any] = {
        "name": c["campaign_name"],
        "objective": objective,
        "status": "PAUSED",
        "special_ad_categories": special,   # API requires the field, even if empty
        "buying_type": (c.get("buying_type") or "AUCTION").upper(),
    }
    if daily:
        campaign_payload["daily_budget"] = daily
    if lifetime:
        campaign_payload["lifetime_budget"] = lifetime
    bid_strategy = (c.get("bid_strategy") or "").upper()
    if cbo:
        campaign_payload["bid_strategy"] = bid_strategy or "LOWEST_COST_WITHOUT_CAP"
    campaign_payload["_page_id"] = c["page_id"]
    campaign_payload["_instagram_user_id"] = c.get("instagram_user_id") or c.get("instagram_actor_id") or None
    campaign_payload["_pixel_id"] = c.get("pixel_id") or None
    campaign_payload["_conversion_event"] = (c.get("conversion_event") or "LEAD").upper()
    campaign_payload["_url_tags"] = c.get("url_tags") or None
    campaign_payload["_bid_strategy"] = bid_strategy

    default_country = (c.get("default_country") or "US").upper()
    resolver = (resolver_factory or (lambda a, k: Resolver(a, k)))(acct, default_country)

    # ---- ad sets ---------------------------------------------------------
    if not data["adsets"]:
        raise BuildError("AdSets sheet has no rows.")
    adsets: list[dict[str, Any]] = []
    seen: set[str] = set()
    for i, row in enumerate(data["adsets"], start=2):
        name = row.get("adset_name", "")
        if not name:
            raise BuildError(f"AdSets row {i}: adset_name is required.")
        if name in seen:
            raise BuildError(f"AdSets: duplicate adset_name '{name}'.")
        seen.add(name)
        goal = (row.get("optimization_goal") or "").upper()
        if goal not in _OPT_GOALS:
            raise BuildError(f"AdSets '{name}': optimization_goal '{row.get('optimization_goal')}' invalid.")
        allowed = _GOALS_BY_OBJECTIVE.get(objective, _OPT_GOALS)
        if goal not in allowed:
            raise BuildError(
                f"AdSets '{name}': optimization_goal {goal} isn't available under {objective}. "
                f"Allowed: {', '.join(sorted(allowed))}. (Or change the campaign objective.)"
            )
        payload: dict[str, Any] = {
            "name": name,
            "optimization_goal": goal,
            "billing_event": (row.get("billing_event") or "IMPRESSIONS").upper(),
            "status": "PAUSED",
            "targeting": build_targeting(row, resolver),
        }
        d = _dollars_to_minor(row.get("daily_budget"), f"{name} daily_budget")
        lt = _dollars_to_minor(row.get("lifetime_budget"), f"{name} lifetime_budget")
        if cbo and (d or lt):
            raise BuildError(f"AdSets '{name}': campaign uses CBO — clear the ad-set budget columns.")
        if not cbo and not (d or lt):
            raise BuildError(f"AdSets '{name}': needs daily_budget or lifetime_budget (campaign has no CBO budget).")
        if d:
            payload["daily_budget"] = d
        if lt:
            payload["lifetime_budget"] = lt
            if not row.get("end_time"):
                raise BuildError(f"AdSets '{name}': lifetime_budget requires end_time.")
        if not cbo:
            payload["bid_strategy"] = bid_strategy or "LOWEST_COST_WITHOUT_CAP"
        bid = _dollars_to_minor(row.get("bid_amount"), f"{name} bid_amount")
        if bid:
            payload["bid_amount"] = bid
        eff_strategy = payload.get("bid_strategy") or campaign_payload.get("bid_strategy") or ""
        if eff_strategy in {"LOWEST_COST_WITH_BID_CAP", "COST_CAP"} and not bid:
            raise BuildError(f"AdSets '{name}': {eff_strategy} requires bid_amount.")
        if row.get("start_time"):
            payload["start_time"] = row["start_time"]
        if row.get("end_time"):
            payload["end_time"] = row["end_time"]
        if goal in {"OFFSITE_CONVERSIONS", "LANDING_PAGE_VIEWS", "LINK_CLICKS"} and objective in {
            "OUTCOME_LEADS", "OUTCOME_SALES", "OUTCOME_TRAFFIC", "OUTCOME_ENGAGEMENT"
        }:
            payload["destination_type"] = "WEBSITE"
        if goal == "OFFSITE_CONVERSIONS":
            if not campaign_payload["_pixel_id"]:
                raise BuildError(f"AdSets '{name}': OFFSITE_CONVERSIONS needs pixel_id on the Campaign sheet.")
            payload["promoted_object"] = {
                "pixel_id": campaign_payload["_pixel_id"],
                "custom_event_type": campaign_payload["_conversion_event"],
            }
        if goal == "LEAD_GENERATION":
            payload["promoted_object"] = {"page_id": c["page_id"]}
            payload["destination_type"] = "ON_AD"
            warnings.append(f"{name}: LEAD_GENERATION (instant forms) — ads need a lead form; this builder creates link ads only.")
        adsets.append({"name": name, "payload": payload, "row": row})

    # ---- ads -------------------------------------------------------------
    if not data["ads"]:
        raise BuildError("Ads sheet has no rows.")
    folder = wb_path if wb_path.is_dir() else wb_path.parent
    assets = _find_assets(folder)
    ads: list[dict[str, Any]] = []
    ad_names: set[str] = set()
    adset_names = {a["name"] for a in adsets}
    dco_adsets: set[str] = set()
    for i, row in enumerate(data["ads"], start=2):
        name = row.get("ad_name", "")
        if not name:
            raise BuildError(f"Ads row {i}: ad_name is required.")
        if name in ad_names:
            raise BuildError(f"Ads: duplicate ad_name '{name}'.")
        ad_names.add(name)
        if row.get("adset_name") not in adset_names:
            raise BuildError(f"Ads '{name}': adset_name '{row.get('adset_name')}' is not on the AdSets sheet.")
        fmt = (row.get("format") or "").lower()
        if fmt not in {"image", "video", "carousel", "dco"}:
            raise BuildError(f"Ads '{name}': format must be image | video | carousel | dco.")
        files = _split(row.get("asset_files"))
        if not files:
            raise BuildError(f"Ads '{name}': asset_files is required.")
        for f in files:
            if f not in assets:
                raise BuildError(f"Ads '{name}': asset '{f}' not found in {folder} (searched subfolders too).")
        kinds = {("video" if assets[f].suffix.lower() in _VIDEO_EXT else "image") for f in files}
        if fmt in {"image", "video"}:
            if len(files) != 1:
                raise BuildError(f"Ads '{name}': {fmt} ads take exactly one asset.")
            if kinds != {fmt}:
                raise BuildError(f"Ads '{name}': format is {fmt} but the file is a {kinds.pop()}.")
        if fmt == "carousel" and not 2 <= len(files) <= 10:
            raise BuildError(f"Ads '{name}': carousel needs 2-10 assets.")
        if fmt == "dco":
            if len(files) > 10:
                raise BuildError(f"Ads '{name}': dco allows up to 10 assets.")
            if len(kinds) > 1:
                raise BuildError(f"Ads '{name}': dco assets must be all images or all videos.")
            if row["adset_name"] in dco_adsets:
                raise BuildError(f"Ads '{name}': an ad set may hold only ONE dco ad.")
            dco_adsets.add(row["adset_name"])
        if not row.get("destination_url"):
            raise BuildError(f"Ads '{name}': destination_url is required.")
        if not row.get("primary_text"):
            raise BuildError(f"Ads '{name}': primary_text is required.")
        for cta in _variants(row.get("cta") or "LEARN_MORE"):
            if cta.upper() not in _CTA_ALLOWED:
                raise BuildError(f"Ads '{name}': cta '{cta}' is not a valid CTA type.")
        if len(_variants(row.get("primary_text"))) > 5 or len(_variants(row.get("headline"))) > 5:
            raise BuildError(f"Ads '{name}': max 5 primary_text / headline variants.")
        for h in _variants(row.get("headline")):
            if len(h) > 255:
                raise BuildError(f"Ads '{name}': headline over 255 chars.")
        if fmt != "dco" and len(_variants(row.get("primary_text"))) > 1:
            warnings.append(f"{name}: multiple primary_text variants only apply to dco — first one will be used.")
        ads.append({"name": name, "adset_name": row["adset_name"], "format": fmt, "assets": files, "row": row})

    for a in adsets:
        if a["name"] in dco_adsets:
            a["payload"]["is_dynamic_creative"] = True
            others = [ad["name"] for ad in ads if ad["adset_name"] == a["name"] and ad["format"] != "dco"]
            if others:
                raise BuildError(f"AdSet '{a['name']}' has a dco ad; it cannot also hold standard ads ({', '.join(others)}).")
        if not any(ad["adset_name"] == a["name"] for ad in ads):
            warnings.append(f"Ad set '{a['name']}' has no ads.")

    return Plan(
        workbook=wb_path, folder=folder, ad_account_id=acct, campaign=campaign_payload,
        adsets=adsets, ads=ads, assets={f: assets[f] for ad in ads for f in ad["assets"]},
        notes=resolver.notes, warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Creative payloads
# ---------------------------------------------------------------------------


def _cta(cta_type: str, link: str) -> dict[str, Any]:
    t = cta_type.upper()
    if t == "NO_BUTTON":
        return {"type": t}
    return {"type": t, "value": {"link": link}}


def creative_payload(
    ad: dict[str, Any],
    campaign: dict[str, Any],
    image_hashes: dict[str, str],
    videos: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Build the /adcreatives body for one Ads row.

    image_hashes: asset filename → image hash
    videos:       asset filename → {"id": video_id, "thumbnail_url": ...}
    """
    row = ad["row"]
    fmt = ad["format"]
    files = ad["assets"]
    bodies = _variants(row.get("primary_text"))
    titles = _variants(row.get("headline"))
    descs = _variants(row.get("description"))
    links = _variants(row.get("destination_url"))
    ctas = [c.upper() for c in _variants(row.get("cta") or "LEARN_MORE")]
    url_tags = row.get("url_tags") or campaign.get("_url_tags")
    display = _s(row.get("display_link")) or None

    story: dict[str, Any] = {"page_id": campaign["_page_id"]}
    if campaign.get("_instagram_user_id"):
        # Meta rejected `instagram_actor_id` for a valid IG id on a live account
        # (Sept 2026); `instagram_user_id` is the field its own creatives carry.
        story["instagram_user_id"] = campaign["_instagram_user_id"]

    out: dict[str, Any] = {"name": ad["name"], "status": "PAUSED"}
    if url_tags:
        out["url_tags"] = url_tags

    if fmt == "image":
        ld: dict[str, Any] = {
            "image_hash": image_hashes[files[0]],
            "link": links[0],
            "message": bodies[0],
            "call_to_action": _cta(ctas[0], links[0]),
        }
        if titles:
            ld["name"] = titles[0]
        if descs:
            ld["description"] = descs[0]
        if display:
            ld["caption"] = display
        story["link_data"] = ld
        out["object_story_spec"] = story

    elif fmt == "video":
        v = videos[files[0]]
        vd: dict[str, Any] = {
            "video_id": v["id"],
            "message": bodies[0],
            "call_to_action": _cta(ctas[0], links[0]),
        }
        if v.get("thumbnail_url"):
            vd["image_url"] = v["thumbnail_url"]
        elif v.get("image_hash"):
            vd["image_hash"] = v["image_hash"]
        if titles:
            vd["title"] = titles[0]
        if descs:
            vd["link_description"] = descs[0]
        story["video_data"] = vd
        out["object_story_spec"] = story

    elif fmt == "carousel":
        cards = []
        for idx, f in enumerate(files):
            link = links[idx] if idx < len(links) else links[0]
            card: dict[str, Any] = {"link": link, "call_to_action": _cta(ctas[0], link)}
            if f in videos:
                card["video_id"] = videos[f]["id"]
                if videos[f].get("thumbnail_url"):
                    card["picture"] = videos[f]["thumbnail_url"]
            else:
                card["image_hash"] = image_hashes[f]
            if titles:
                card["name"] = titles[idx] if idx < len(titles) else titles[0]
            if descs:
                card["description"] = descs[idx] if idx < len(descs) else descs[0]
            cards.append(card)
        ld = {
            "link": links[0],
            "message": bodies[0],
            "child_attachments": cards,
            "multi_share_optimized": True,
            "multi_share_end_card": False,
        }
        if display:
            ld["caption"] = display
        story["link_data"] = ld
        out["object_story_spec"] = story

    elif fmt == "dco":
        afs: dict[str, Any] = {
            "bodies": [{"text": b} for b in bodies],
            "link_urls": [{"website_url": links[0], **({"display_url": display} if display else {})}],
            "call_to_action_types": ctas,
        }
        if titles:
            afs["titles"] = [{"text": t} for t in titles]
        if descs:
            afs["descriptions"] = [{"text": d} for d in descs]
        if files[0] in videos:
            afs["videos"] = [{"video_id": videos[f]["id"], **({"thumbnail_url": videos[f]["thumbnail_url"]} if videos[f].get("thumbnail_url") else {})} for f in files]
            afs["ad_formats"] = ["SINGLE_VIDEO"]
        else:
            afs["images"] = [{"hash": image_hashes[f]} for f in files]
            afs["ad_formats"] = ["SINGLE_IMAGE"]
        out["object_story_spec"] = story
        out["asset_feed_spec"] = afs

    return out


def ad_payload(ad: dict[str, Any], adset_id: str, creative_id: str, campaign: dict[str, Any]) -> dict[str, Any]:
    p: dict[str, Any] = {
        "name": ad["name"],
        "adset_id": adset_id,
        "creative": {"creative_id": creative_id},
        "status": "PAUSED",
    }
    if campaign.get("_pixel_id"):
        p["tracking_specs"] = [{"action.type": ["offsite_conversion"], "fb_pixel": [campaign["_pixel_id"]]}]
    return p


# ---------------------------------------------------------------------------
# Execute (resumable)
# ---------------------------------------------------------------------------


def _state_path(plan_: Plan) -> Path:
    if plan_.workbook.is_dir() or plan_.workbook.suffix.lower() == ".csv":
        return plan_.folder / "campaign.build-state.json"
    return plan_.workbook.with_suffix(".build-state.json")


def _results_path(plan_: Plan) -> Path:
    if plan_.workbook.is_dir() or plan_.workbook.suffix.lower() == ".csv":
        return plan_.folder / "campaign-build-results.csv"
    return plan_.workbook.with_name(f"{plan_.workbook.stem}-build-results.csv")


def _load_state(plan_: Plan) -> dict[str, Any]:
    p = _state_path(plan_)
    if p.is_file():
        try:
            st = json.loads(p.read_text())
            if st.get("campaign_name") == plan_.campaign["name"] and st.get("ad_account_id") == plan_.ad_account_id:
                return st
        except json.JSONDecodeError:
            pass
    return {
        "campaign_name": plan_.campaign["name"], "ad_account_id": plan_.ad_account_id,
        "images": {}, "videos": {}, "campaign_id": None, "adsets": {}, "creatives": {}, "ads": {},
    }


def _save_state(plan_: Plan, st: dict[str, Any]) -> None:
    _state_path(plan_).write_text(json.dumps(st, indent=2))


def execute(plan_: Plan, progress: Callable[[str], None] | None = None) -> dict[str, Any]:
    """Create everything in the plan, PAUSED. Returns IDs + the results CSV path."""
    log = progress or (lambda m: None)
    st = _load_state(plan_)
    acct = plan_.ad_account_id
    camp = plan_.campaign

    # 1. assets
    for fname, path in plan_.assets.items():
        if path.suffix.lower() in _VIDEO_EXT:
            if fname not in st["videos"]:
                log(f"uploading video {fname}")
                st["videos"][fname] = graph.upload_video(path, acct, title=fname)
                _save_state(plan_, st)
        else:
            if fname not in st["images"]:
                log(f"uploading image {fname}")
                st["images"][fname] = graph.upload_image(path, acct)["hash"]
                _save_state(plan_, st)

    # 2. campaign
    if not st["campaign_id"]:
        body = {k: v for k, v in camp.items() if not k.startswith("_")}
        log(f"creating campaign {body['name']}")
        st["campaign_id"] = graph.graph_post(f"{acct}/campaigns", body)["id"]
        _save_state(plan_, st)
    campaign_id = st["campaign_id"]

    # 3. ad sets
    for a in plan_.adsets:
        if a["name"] in st["adsets"]:
            continue
        body = dict(a["payload"], campaign_id=campaign_id)
        log(f"creating ad set {a['name']}")
        st["adsets"][a["name"]] = graph.graph_post(f"{acct}/adsets", body)["id"]
        _save_state(plan_, st)

    # 4. creatives + ads
    for ad in plan_.ads:
        if ad["name"] not in st["creatives"]:
            body = creative_payload(ad, camp, st["images"], st["videos"])
            log(f"creating creative {ad['name']}")
            st["creatives"][ad["name"]] = graph.graph_post(f"{acct}/adcreatives", body)["id"]
            _save_state(plan_, st)
        if ad["name"] not in st["ads"]:
            body = ad_payload(ad, st["adsets"][ad["adset_name"]], st["creatives"][ad["name"]], camp)
            log(f"creating ad {ad['name']}")
            st["ads"][ad["name"]] = graph.graph_post(f"{acct}/ads", body)["id"]
            _save_state(plan_, st)

    # 5. results file
    results = _results_path(plan_)
    with results.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["ad_name", "adset_name", "format", "campaign_id", "adset_id", "creative_id", "ad_id", "status"])
        for ad in plan_.ads:
            w.writerow([
                ad["name"], ad["adset_name"], ad["format"], campaign_id,
                st["adsets"][ad["adset_name"]], st["creatives"][ad["name"]], st["ads"][ad["name"]], "PAUSED",
            ])

    return {
        "campaign_id": campaign_id,
        "adsets": st["adsets"],
        "ads": st["ads"],
        "creatives": st["creatives"],
        "results_csv": str(results),
        "state_file": str(_state_path(plan_)),
        "ads_manager_url": graph.ad_preview_url(acct, campaign_id),
    }


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------


def write_template(path: str | Path) -> Path:
    """Write a blank campaign.xlsx with headers, help rows and one example per sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = Path(path).expanduser()
    if p.is_dir():
        p = p / "campaign.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")
    help_font = Font(italic=True, color="808080")

    ws = wb.active
    ws.title = "Campaign"
    ws.append(["key", "value", "help"])
    for k, h in CAMPAIGN_KEYS.items():
        ws.append([k, "", h])
    ws["B2"] = "Acme Equipment"
    ws["B3"] = "act_"
    ws["B6"] = "AE | Fall Promo | Meta | BOFU"
    ws["B7"] = "OUTCOME_LEADS"
    ws["B8"] = "AUCTION"
    ws["B10"] = "150.00"
    ws["B12"] = "LOWEST_COST_WITHOUT_CAP"
    ws["B14"] = "LEAD"
    ws["B15"] = "US"
    ws["B16"] = "utm_source=facebook&utm_medium=ppc&utm_campaign={{campaign.name}}&utm_content={{ad.name}}"
    for c in ws[1]:
        c.font, c.fill = head, fill
    for row in ws.iter_rows(min_row=2):
        row[2].font = help_font
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 110

    def sheet(title: str, cols: dict[str, str], example: list[str]) -> None:
        w = wb.create_sheet(title)
        w.append(list(cols))
        w.append(["# " + h for h in cols.values()])
        w.append(example)
        for c in w[1]:
            c.font, c.fill = head, fill
        for c in w[2]:
            c.font = help_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        w.row_dimensions[2].height = 90
        for i, _ in enumerate(cols, start=1):
            w.column_dimensions[get_column_letter(i)].width = 28
        w.freeze_panes = "A3"

    sheet("AdSets", ADSET_COLUMNS, [
        "Radius | Nashville dealers", "", "", "OFFSITE_CONVERSIONS", "IMPRESSIONS", "",
        "radius", "1200 Broadway, Nashville, TN 37203 | 25mi; 100 Main St, Franklin, TN 37064 | 20mi", "", "",
        "25", "65", "all", "Carpet cleaning; Small business owners", "", "", "Y", "auto", "", "",
    ])
    sheet("Ads", AD_COLUMNS, [
        "Hero | Day One Kit | Static", "Radius | Nashville dealers", "image", "hero-day-one-kit.png",
        "Buy any new pressure washer this fall and we ship the Day One Kit free.",
        "Free Day One Kit With Any Pressure Washer", "Offer ends October 31",
        "https://www.example.com/fall-promo", "example.com", "GET_QUOTE", "",
    ])
    wb.save(p)
    return p
