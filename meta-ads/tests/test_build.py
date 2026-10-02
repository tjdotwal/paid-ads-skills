"""Unit tests for the spreadsheet campaign builder. No network: every Graph
call is intercepted and recorded so we can assert on the exact payloads.

Run:  python3 -m pytest tests/  (or python3 -m unittest tests.test_build)
"""

from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from mcp_meta_ads import build, graph



def _png(path: Path, w: int = 64, h: int = 64) -> None:
    """Write a tiny valid PNG so asset discovery + upload paths have real files."""
    import struct
    import zlib

    raw = b"".join(b"\x00" + b"\x1f\x3a\x5f" * w for _ in range(h))

    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def make_fixtures(folder: Path) -> None:
    (folder / "assets").mkdir(exist_ok=True)
    _png(folder / "hero-day-one-kit.png")
    for n in ("card1", "card2", "card3"):
        _png(folder / "assets" / f"{n}.png")
    (folder / "assets" / "testimonial.mp4").write_bytes(b"\x00" * 4096)


class FakeGraph:
    """Stand-in for graph.* network functions. Records POSTs, answers lookups."""

    def __init__(self):
        self.posts: list[tuple[str, dict]] = []
        self.counter = 1000

    def _id(self):
        self.counter += 1
        return str(self.counter)

    def graph_post(self, path, data=None, files=None, timeout=120):
        self.posts.append((path, json.loads(json.dumps(data or {}))))
        return {"id": self._id()}

    def upload_image(self, path, ad_account_id=None):
        return {"hash": f"hash_{Path(path).stem}", "url": "https://cdn/x.png", "name": Path(path).name}

    def upload_video(self, path, ad_account_id=None, title=None, wait_ready=True, timeout_s=600):
        return {"id": f"vid_{Path(path).stem}", "status": "ready", "thumbnail_url": "https://cdn/thumb.jpg"}

    def search_geo(self, q, location_types, country_code=None, region_id=None, limit=25):
        if "city" in location_types:
            return [
                {"key": "2490299", "name": "Nashville", "region": "Tennessee", "country_code": "US"},
                {"key": "2490300", "name": "Nashville", "region": "Illinois", "country_code": "US"},
                {"key": "2481000", "name": "Franklin", "region": "Tennessee", "country_code": "US"},
            ] if q in ("Nashville", "Franklin") else []
        if "zip" in location_types:
            return [{"key": f"US:{q}", "name": q, "primary_city": "Nashville"}]
        if "region" in location_types:
            return [{"key": "3874", "name": "Tennessee"}]
        return []

    def search_interests(self, q, limit=25):
        table = {
            "carpet cleaning": [{"id": "6003100", "name": "Carpet cleaning", "audience_size_lower_bound": 1200000}],
            "small business owners": [{"id": "6003200", "name": "Small business", "audience_size_lower_bound": 90000000}],
        }
        return table.get(q.lower(), [])

    def list_custom_audiences(self, ad_account_id=None):
        return [
            {"id": "23850000001", "name": "Website visitors 180d"},
            {"id": "23850000002", "name": "Customers - all time"},
        ]


def install(fake: FakeGraph):
    for name in ("graph_post", "upload_image", "upload_video", "search_geo", "search_interests", "list_custom_audiences"):
        setattr(graph, name, getattr(fake, name))


def fill_workbook(path: Path, *, cbo=True):
    wb = load_workbook(path)
    c = wb["Campaign"]
    vals = {r[0].value: r for r in c.iter_rows(min_row=2)}
    def setv(k, v): vals[k][1].value = v
    setv("ad_account_id", "act_123456789")
    setv("page_id", "111222333")
    setv("campaign_name", "AE | Fall Promo | Meta | BOFU")
    setv("objective", "OUTCOME_LEADS")
    setv("campaign_daily_budget", "150.00" if cbo else "")
    setv("pixel_id", "987654321")
    setv("conversion_event", "LEAD")
    setv("url_tags", "utm_source=facebook&utm_medium=ppc&utm_campaign={{campaign.name}}&utm_content={{ad.name}}")

    a = wb["AdSets"]
    a.delete_rows(3, a.max_row)
    cols = list(build.ADSET_COLUMNS)
    def adset(**kw):
        a.append([kw.get(k, "") for k in cols])
    adset(adset_name="Radius | Nashville dealers", optimization_goal="OFFSITE_CONVERSIONS",
          geo_type="radius", geo_value="1200 Broadway, Nashville, TN 37203 | 25mi; 36.16,-86.78 | 10mi",
          age_min="25", age_max="65", interests="Carpet cleaning; Small business owners",
          excluded_audiences="Customers - all time", daily_budget="" if cbo else "60")
    adset(adset_name="Cities | TN metros", optimization_goal="OFFSITE_CONVERSIONS",
          geo_type="cities", geo_value="Nashville, TN | 15mi; Franklin, TN", genders="male",
          custom_audiences="Website visitors 180d", advantage_audience="N",
          placements="facebook:feed,story; instagram:stream,reels", daily_budget="" if cbo else "40")
    adset(adset_name="Zips | DCO", optimization_goal="OFFSITE_CONVERSIONS", geo_type="zips",
          geo_value="37203; 37204", daily_budget="" if cbo else "30")

    d = wb["Ads"]
    d.delete_rows(3, d.max_row)
    acols = list(build.AD_COLUMNS)
    def ad(**kw):
        d.append([kw.get(k, "") for k in acols])
    ad(ad_name="Hero | Static", adset_name="Radius | Nashville dealers", format="image", asset_files="hero-day-one-kit.png",
       primary_text="Buy any new pressure washer this fall and we ship the Day One Kit free.",
       headline="Free Day One Kit", description="Offer ends October 31",
       destination_url="https://www.example.com/fall-promo", display_link="example.com", cta="GET_QUOTE")
    ad(ad_name="Testimonial | Video", adset_name="Radius | Nashville dealers", format="video", asset_files="testimonial.mp4",
       primary_text="25 years in business, best machine I have owned.", headline="Hear it from Brian",
       destination_url="https://www.example.com/fall-promo", cta="LEARN_MORE")
    ad(ad_name="Lineup | Carousel", adset_name="Cities | TN metros", format="carousel", asset_files="card1.png; card2.png; card3.png",
       primary_text="Three pressure washers. One free Day One Kit.", headline="Model 370 || Model 570 || Model 870",
       description="From $29,729 || From $37,392 || From $49,313",
       destination_url="https://www.example.com/fall-promo#370 || https://www.example.com/fall-promo#570 || https://www.example.com/fall-promo#870",
       cta="GET_QUOTE", url_tags="utm_source=facebook&utm_medium=ppc&utm_content=carousel")
    ad(ad_name="DCO | Mix", adset_name="Zips | DCO", format="dco", asset_files="card1.png; card2.png",
       primary_text="Copy A || Copy B", headline="Head 1 || Head 2",
       destination_url="https://www.example.com/fall-promo", cta="GET_QUOTE || LEARN_MORE")
    wb.save(path)


class BuildTests(unittest.TestCase):
    def setUp(self):
        self.fake = FakeGraph()
        install(self.fake)
        self.dir = Path(tempfile.mkdtemp())
        make_fixtures(self.dir)
        self.wb = build.write_template(self.dir)
        fill_workbook(self.wb)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_template_has_three_sheets(self):
        wb = load_workbook(self.wb)
        self.assertEqual([s.lower() for s in wb.sheetnames], ["campaign", "adsets", "ads"])

    def test_plan_is_side_effect_free(self):
        p = build.plan(self.wb)
        self.assertEqual(self.fake.posts, [])
        self.assertEqual(len(p.adsets), 3)
        self.assertEqual(len(p.ads), 4)
        self.assertIn("Everything will be created PAUSED", p.summary())

    def test_campaign_payload(self):
        p = build.plan(self.wb)
        c = {k: v for k, v in p.campaign.items() if not k.startswith("_")}
        self.assertEqual(c["daily_budget"], 15000)
        self.assertEqual(c["special_ad_categories"], [])
        self.assertEqual(c["status"], "PAUSED")
        self.assertEqual(c["bid_strategy"], "LOWEST_COST_WITHOUT_CAP")

    def test_radius_targeting(self):
        p = build.plan(self.wb)
        t = p.adsets[0]["payload"]["targeting"]
        locs = t["geo_locations"]["custom_locations"]
        self.assertEqual(locs[0], {"radius": 25.0, "distance_unit": "mile", "address_string": "1200 Broadway, Nashville, TN 37203"})
        self.assertEqual(locs[1]["latitude"], 36.16)
        self.assertEqual(t["age_min"], 25)
        self.assertEqual(t["flexible_spec"], [{"interests": [
            {"id": "6003100", "name": "Carpet cleaning"}, {"id": "6003200", "name": "Small business"}]}])
        self.assertEqual(t["excluded_custom_audiences"], [{"id": "23850000002"}])
        # interests present → Advantage+ audience defaults OFF unless told otherwise
        self.assertEqual(t["targeting_automation"], {"advantage_audience": 0})
        self.assertIn("nearest match 'Small business'", " ".join(p.notes))
        ps = p.adsets[0]["payload"]
        self.assertEqual(ps["promoted_object"], {"pixel_id": "987654321", "custom_event_type": "LEAD"})
        self.assertEqual(ps["destination_type"], "WEBSITE")
        self.assertNotIn("daily_budget", ps)  # CBO
        self.assertNotIn("bid_strategy", ps)

    def test_city_targeting_and_placements(self):
        p = build.plan(self.wb)
        t = p.adsets[1]["payload"]["targeting"]
        self.assertEqual(t["geo_locations"]["cities"], [
            {"key": "2490299", "radius": 15.0, "distance_unit": "mile"}, {"key": "2481000"}])
        self.assertEqual(t["genders"], [1])
        self.assertEqual(t["custom_audiences"], [{"id": "23850000001"}])
        self.assertEqual(t["publisher_platforms"], ["facebook", "instagram"])
        self.assertEqual(t["facebook_positions"], ["feed", "story"])
        self.assertEqual(t["instagram_positions"], ["stream", "reels"])

    def test_zip_targeting_and_dco_flag(self):
        p = build.plan(self.wb)
        ps = p.adsets[2]["payload"]
        self.assertEqual(ps["targeting"]["geo_locations"]["zips"], [{"key": "US:37203"}, {"key": "US:37204"}])
        self.assertTrue(ps["is_dynamic_creative"])
        self.assertEqual(ps["targeting"]["targeting_automation"], {"advantage_audience": 1})

    def test_execute_creates_everything_paused_and_is_resumable(self):
        p = build.plan(self.wb)
        log: list[str] = []
        res = build.execute(p, log.append)
        paths = [pp for pp, _ in self.fake.posts]
        self.assertEqual(paths.count("act_123456789/campaigns"), 1)
        self.assertEqual(paths.count("act_123456789/adsets"), 3)
        self.assertEqual(paths.count("act_123456789/adcreatives"), 4)
        self.assertEqual(paths.count("act_123456789/ads"), 4)
        for _, body in self.fake.posts:
            self.assertEqual(body.get("status"), "PAUSED")
        self.assertTrue(Path(res["results_csv"]).is_file())
        self.assertTrue(Path(res["state_file"]).is_file())
        self.assertIn("selected_campaign_ids=", res["ads_manager_url"])

        # Rerun: nothing new is posted.
        before = len(self.fake.posts)
        build.execute(build.plan(self.wb))
        self.assertEqual(len(self.fake.posts), before)

    def test_creative_payloads(self):
        p = build.plan(self.wb)
        build.execute(p)
        creatives = {b["name"]: b for pp, b in self.fake.posts if pp.endswith("/adcreatives")}

        img = creatives["Hero | Static"]
        ld = img["object_story_spec"]["link_data"]
        self.assertEqual(ld["image_hash"], "hash_hero-day-one-kit")
        self.assertEqual(ld["name"], "Free Day One Kit")
        self.assertEqual(ld["caption"], "example.com")
        self.assertEqual(ld["call_to_action"], {"type": "GET_QUOTE", "value": {"link": "https://www.example.com/fall-promo"}})
        self.assertIn("utm_campaign={{campaign.name}}", img["url_tags"])

        vid = creatives["Testimonial | Video"]["object_story_spec"]["video_data"]
        self.assertEqual(vid["video_id"], "vid_testimonial")
        self.assertEqual(vid["image_url"], "https://cdn/thumb.jpg")
        self.assertEqual(vid["title"], "Hear it from Brian")

        car = creatives["Lineup | Carousel"]
        cards = car["object_story_spec"]["link_data"]["child_attachments"]
        self.assertEqual(len(cards), 3)
        self.assertEqual(cards[1]["name"], "Model 570")
        self.assertEqual(cards[2]["link"], "https://www.example.com/fall-promo#870")
        self.assertEqual(car["url_tags"], "utm_source=facebook&utm_medium=ppc&utm_content=carousel")  # per-ad override

        dco = creatives["DCO | Mix"]
        afs = dco["asset_feed_spec"]
        self.assertEqual([b["text"] for b in afs["bodies"]], ["Copy A", "Copy B"])
        self.assertEqual(afs["call_to_action_types"], ["GET_QUOTE", "LEARN_MORE"])
        self.assertEqual(afs["ad_formats"], ["SINGLE_IMAGE"])
        self.assertEqual(len(afs["images"]), 2)

        ads = [b for pp, b in self.fake.posts if pp.endswith("/ads")]
        self.assertEqual(ads[0]["tracking_specs"], [{"action.type": ["offsite_conversion"], "fb_pixel": ["987654321"]}])

    def test_validation_errors_are_specific(self):
        wb = load_workbook(self.wb)
        ws = wb["Ads"]
        ws.cell(row=3, column=list(build.AD_COLUMNS).index("asset_files") + 1).value = "missing.png"
        wb.save(self.wb)
        with self.assertRaises(build.BuildError) as cm:
            build.plan(self.wb)
        self.assertIn("missing.png", str(cm.exception))

    def test_cbo_conflict_detected(self):
        wb = load_workbook(self.wb)
        ws = wb["AdSets"]
        ws.cell(row=3, column=list(build.ADSET_COLUMNS).index("daily_budget") + 1).value = "50"
        wb.save(self.wb)
        with self.assertRaises(build.BuildError) as cm:
            build.plan(self.wb)
        self.assertIn("CBO", str(cm.exception))

    def test_adset_budgets_without_cbo(self):
        fill_workbook(self.wb, cbo=False)
        p = build.plan(self.wb)
        self.assertNotIn("daily_budget", {k for k in p.campaign if not k.startswith("_")})
        self.assertEqual(p.adsets[0]["payload"]["daily_budget"], 6000)
        self.assertEqual(p.adsets[0]["payload"]["bid_strategy"], "LOWEST_COST_WITHOUT_CAP")

    def test_google_sheets_csv_export_is_equivalent(self):
        """Export each tab of the filled workbook as '<Sheet> - <Tab>.csv' the way
        Google Sheets names them, delete the xlsx, and build from the folder."""
        import csv as _csv
        wb = load_workbook(self.wb, data_only=True)
        for tab in ("Campaign", "AdSets", "Ads"):
            with (self.dir / f"Acme Fall Promo - {tab}.csv").open("w", newline="", encoding="utf-8") as fh:
                w = _csv.writer(fh)
                for row in wb[tab].iter_rows(values_only=True):
                    w.writerow(["" if c is None else c for c in row])
        xlsx_plan = build.plan(self.wb)
        self.wb.unlink()
        csv_plan = build.plan(self.dir)  # folder path
        self.assertEqual([a["payload"] for a in csv_plan.adsets], [a["payload"] for a in xlsx_plan.adsets])
        self.assertEqual({k: v for k, v in csv_plan.campaign.items()}, {k: v for k, v in xlsx_plan.campaign.items()})
        self.assertEqual([a["name"] for a in csv_plan.ads], [a["name"] for a in xlsx_plan.ads])
        # Also accepts a path to any one of the CSVs.
        one = build.plan(self.dir / "Acme Fall Promo - Ads.csv")
        self.assertEqual(len(one.ads), 4)
        res = build.execute(one)
        self.assertEqual(Path(res["results_csv"]).name, "campaign-build-results.csv")
        self.assertEqual(Path(res["state_file"]).parent.resolve(), self.dir.resolve())

    def test_live_lessons_from_first_live_run(self):
        """Rules learned from the first live build on act_1484941429463480."""
        r = build.Resolver("act_1", "US")
        # 1. Advantage+ audience forces age_max to 65 (Meta error 1870189).
        t = build.build_targeting({"geo_type": "countries", "geo_value": "US", "age_min": "18", "age_max": "45",
                                   "advantage_audience": "Y"}, r)
        self.assertEqual(t["age_max"], 65)
        self.assertTrue(any("age_max 45 → 65" in n for n in r.notes))
        # ...but a hard ceiling is kept when Advantage+ is off.
        t = build.build_targeting({"geo_type": "countries", "geo_value": "US", "age_max": "45", "advantage_audience": "N"}, r)
        self.assertEqual(t["age_max"], 45)
        # 2. Bare city names can't be geocoded (1487855) — rejected up front.
        with self.assertRaises(build.BuildError) as cm:
            build.build_targeting({"geo_type": "radius", "geo_value": "Calgary, AB, Canada | 25mi"}, r)
        self.assertIn("street number", str(cm.exception))
        # 3. Goal must fit the objective (2490408).
        wb = load_workbook(self.wb)
        ws = wb["AdSets"]
        ws.cell(row=3, column=list(build.ADSET_COLUMNS).index("optimization_goal") + 1).value = "LANDING_PAGE_VIEWS"
        wb.save(self.wb)
        with self.assertRaises(build.BuildError) as cm:
            build.plan(self.wb)
        self.assertIn("isn't available under OUTCOME_LEADS", str(cm.exception))
        # 4. Instagram identity goes out as instagram_user_id.
        ad = {"name": "x", "format": "image", "assets": ["a.png"], "row": {"primary_text": "t", "destination_url": "https://x.com", "cta": "LEARN_MORE"}}
        body = build.creative_payload(ad, {"_page_id": "1", "_instagram_user_id": "178", "_url_tags": None}, {"a.png": "h"}, {})
        self.assertEqual(body["object_story_spec"]["instagram_user_id"], "178")
        self.assertNotIn("instagram_actor_id", body["object_story_spec"])

    def test_ambiguous_city_without_state_falls_back_to_first(self):
        r = build.Resolver("act_1", "US")
        self.assertEqual(r.city("Nashville")["key"], "2490299")
        self.assertEqual(r.city("Nashville, IL")["key"], "2490300")
        with self.assertRaises(build.BuildError):
            r.city("Nashville, TX")


if __name__ == "__main__":
    unittest.main()
